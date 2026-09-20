#!/usr/bin/env python3
"""
Grade the pools ourselves from ESPN final scores, rather than waiting on the
commissioner's answer key.

Mike fills in each week's results by hand and mails them with the NEXT week's
picks, so his grading of Week N arrives around Sunday of Week N+1. The scores
are final on Tuesday. This closes that gap, and then turns his copy into a
check on ours instead of the only source.

    python scripts/grade_week.py 2026 [week]
    python scripts/grade_week.py 2026 1 --check "Weekly picks 26.xlsx" --check-suicide "Suicide 26.xlsx"

Reads (all committed):
    data/schedule-<year>.json          final scores — run fetch_schedule.py first
    data/number-map-<year>.json        which games count, and each team's number
    data/raw/entries-<year>-w<NN>.json every card, from parse_pool_picks.py
    data/survivor-<year>.json          every suicide pick, from parse_survivor.py

Writes:
    data/results/pickem-<year>-w<NN>.json   answer key, every card graded, the winner
    data/results/survivor-<year>.json       each suicide pick won / lost, per week

A week with games still to play is written with "complete": false and no
ranking, so a Sunday-night run is honest about what it does not know yet.
Files carry no timestamp and are only rewritten when their content changes, so
the Tuesday workflow run commits nothing on a week that did not move.

--check compares a later workbook from Mike against what we graded:
    * his answer key vs. ours (a disagreement is a keying error on one side)
    * cards whose picks or tiebreaker changed since we parsed them — the
      "adjusted from keying errors" corrections — plus entries added or removed
    * his per-card week totals vs. ours
It reports only; it never rewrites data/raw/. Re-run parse_pool_picks.py for
that week if his corrections should become the record.

--check-suicide matches his rows to ours by (nickname, real name), never by
entry number, because his graded sheet renumbers the eliminated entries. It
lists changed picks and retyped nicknames, then compares his red-filled "out"
cells with the losers we graded, for every week that is final. It reports
only; it never rewrites data/survivor-<year>.json.
"""

import argparse
import json
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"


def load(path):
    return json.loads(path.read_text(encoding="utf-8"))


def write_if_changed(path, obj):
    text = json.dumps(obj, indent=1, ensure_ascii=False) + "\n"
    if path.exists() and path.read_text(encoding="utf-8") == text:
        return False
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")
    return True


def mascot_to_abbr():
    teams = load(DATA / "teams" / "team-identity.json")["teams"]
    return {t["mascot"]: abbr for abbr, t in teams.items()}


def outcome(game):
    """-> 'away' | 'home' | 'tie' | None (not final)."""
    if not game or not game.get("completed"):
        return None
    a, h = game["awayScore"], game["homeScore"]
    if a is None or h is None:
        return None
    return "away" if a > h else "home" if h > a else "tie"


def schedule_index(year):
    games = load(DATA / f"schedule-{year}.json")["games"]
    by_pair = {(g["week"], g["away"], g["home"]): g for g in games}
    by_team = {}
    for g in games:
        by_team[(g["week"], g["away"])] = g
        by_team[(g["week"], g["home"])] = g
    return by_pair, by_team


# --------------------------------------------------------------------------- pick'em

def grade_pickem(year, week, by_pair):
    raw_path = DATA / "raw" / f"entries-{year}-w{week:02d}.json"
    raw = load(raw_path)
    abbr = mascot_to_abbr()
    games = [g for g in load(DATA / f"number-map-{year}.json")["weeks"][str(week)] if g["counts"]]

    key, pending, ties, tiebreak_total = {}, [], [], None
    for g in games:
        sched = by_pair.get((week, abbr[g["away"]], abbr[g["home"]]))
        if sched is None:
            sys.exit(f"week {week}: {g['away']} at {g['home']} is not in the schedule feed")
        res = outcome(sched)
        label = f"{g['away']} at {g['home']}"
        if res is None:
            pending.append(label)
        elif res == "tie":
            # The pool's rule for a tie is not on record. Nobody is credited
            # rather than guessing; say so loudly so it gets asked.
            ties.append(label)
        else:
            key[str(g["awayNum"])] = {
                "winner": g[res], "number": g[f"{res}Num"],
                "score": f"{sched['awayScore']}-{sched['homeScore']}",
            }
        if g["tiebreaker"] and res is not None:
            tiebreak_total = sched["awayScore"] + sched["homeScore"]

    winning = {k["number"] for k in key.values()}
    complete = not pending

    graded = []
    for e in raw["entries"]:
        mnf = e["mnf"] if isinstance(e["mnf"], (int, float)) else None
        graded.append({
            "entry": e["entry"], "name": e["name"], "nick": e["nick"],
            "correct": sum(1 for p in e["picks"] if p in winning),
            "picks": len(e["picks"]),
            "mnf": mnf,
            "mnfDiff": abs(mnf - tiebreak_total) if mnf is not None and tiebreak_total is not None else None,
        })

    winners = []
    if complete:
        # Most correct, then closest Monday-night total by absolute value
        # (confirmed rule — see CLAUDE.md "The tiebreaker guess"). A card with
        # no guess sorts behind every card that made one.
        graded.sort(key=lambda r: (-r["correct"], r["mnfDiff"] if r["mnfDiff"] is not None else 10**6))
        rank = 0
        for i, r in enumerate(graded):
            if i == 0 or (r["correct"], r["mnfDiff"]) != (graded[i - 1]["correct"], graded[i - 1]["mnfDiff"]):
                rank = i + 1
            r["rank"] = rank
        winners = [r["entry"] for r in graded if r["rank"] == 1]
    else:
        graded.sort(key=lambda r: (-r["correct"], r["entry"]))

    out = {
        "year": year, "week": week, "source": "espn",
        "complete": complete, "gamesScored": len(games), "gamesFinal": len(games) - len(pending),
        "pending": pending, "ties": ties,
        "tiebreakerTotal": tiebreak_total,
        "answerKey": key,
        "winners": winners,
        "entries": graded,
    }
    changed = write_if_changed(DATA / "results" / f"pickem-{year}-w{week:02d}.json", out)

    print(f"pick'em week {week}: {out['gamesFinal']}/{len(games)} games final"
          + ("" if complete else f" — waiting on {', '.join(pending)}"))
    if ties:
        print(f"  TIE (nobody credited — confirm the pool rule): {', '.join(ties)}")
    if complete:
        top = graded[0]
        who = ", ".join(f"#{r['entry']} {r['nick'] or r['name']}" for r in graded if r["rank"] == 1)
        print(f"  winner: {who} — {top['correct']}/{len(games)}, MNF guess {top['mnf']} vs {tiebreak_total}")
    print(f"  {'wrote' if changed else 'unchanged'} data/results/pickem-{year}-w{week:02d}.json")
    return out


# --------------------------------------------------------------------------- survivor

def grade_survivor(year, by_team):
    path = DATA / f"survivor-{year}.json"
    if not path.exists():
        return None
    surv = load(path)

    weeks = {}
    for e in surv["entries"]:
        for wk, team in e["picks"].items():
            res = outcome(by_team.get((int(wk), team)))
            if res is None:
                result = "pending"
            elif res == "tie":
                result = "tie"
            else:
                sched = by_team[(int(wk), team)]
                result = "won" if sched[res] == team else "lost"
            w = weeks.setdefault(wk, {"won": 0, "lost": 0, "tie": 0, "pending": 0, "losers": [], "tied": []})
            w[result] += 1
            if result in ("lost", "tie"):
                w["losers" if result == "lost" else "tied"].append(
                    {"entry": e["entry"], "nick": e["nick"], "name": e["name"], "team": team})

    # Losers only. "Eliminated" is deliberately not derived: buy-backs exist
    # and cannot be read from any feed (CLAUDE.md, Planning tab "Not built").
    out = {"year": year, "source": "espn", "weeks": weeks}
    changed = write_if_changed(DATA / "results" / f"survivor-{year}.json", out)
    for wk in sorted(weeks, key=int):
        w = weeks[wk]
        by_team_lost = {}
        for l in w["losers"]:
            by_team_lost[l["team"]] = by_team_lost.get(l["team"], 0) + 1
        lost = ", ".join(f"{t} {n}" for t, n in sorted(by_team_lost.items(), key=lambda x: -x[1]))
        print(f"suicide week {wk}: {w['won']} won, {w['lost']} lost"
              + (f", {w['tie']} tied" if w["tie"] else "")
              + (f", {w['pending']} pending" if w["pending"] else "")
              + (f"  (lost: {lost})" if lost else ""))
    print(f"  {'wrote' if changed else 'unchanged'} data/results/survivor-{year}.json")
    return out


# --------------------------------------------------------------------------- checks

def check_pickem(year, week, ours, workbook):
    sys.path.insert(0, str(ROOT / "scripts"))
    import openpyxl
    from parse_pool_picks import read_week, PICK_COLS

    wb = openpyxl.load_workbook(workbook, data_only=True)
    if str(week) not in wb.sheetnames:
        sys.exit(f"{workbook} has no sheet '{week}'")
    games = [g for g in load(DATA / f"number-map-{year}.json")["weeks"][str(week)] if g["counts"]]
    his_key, his_entries = read_week(wb[str(week)], len(games))
    raw = load(DATA / "raw" / f"entries-{year}-w{week:02d}.json")

    print(f"\n=== check: Mike's week {week} vs ours ===")
    problems = 0

    our_nums = {v["number"] for v in ours["answerKey"].values()}
    his_nums = set(his_key.values())
    if not his_key:
        print("answer key: his is blank — nothing to compare")
    elif his_nums == our_nums:
        print(f"answer key: agrees on all {len(our_nums)} games")
    else:
        problems += 1
        print(f"answer key DISAGREES — his only: {sorted(his_nums - our_nums)}, ours only: {sorted(our_nums - his_nums)}")

    before = {e["entry"]: e for e in raw["entries"]}
    after = {e["entry"]: e for e in his_entries}
    for n in sorted(after.keys() - before.keys()):
        problems += 1
        print(f"  added   #{n} {after[n]['nick']}: {after[n]['picks']}")
    for n in sorted(before.keys() - after.keys()):
        problems += 1
        print(f"  removed #{n} {before[n]['nick']}")
    for n in sorted(before.keys() & after.keys()):
        b, a = before[n], after[n]
        if b["picks"] != a["picks"] or b["mnf"] != a["mnf"]:
            problems += 1
            dropped = sorted(set(b["picks"]) - set(a["picks"]))
            added = sorted(set(a["picks"]) - set(b["picks"]))
            mnf = f", MNF {b['mnf']} -> {a['mnf']}" if b["mnf"] != a["mnf"] else ""
            print(f"  changed #{n} {a['nick']}: -{dropped} +{added}{mnf}")

    # His totals vs. ours, graded on HIS cards so a keying fix is not counted twice.
    our_correct = {e["entry"]: e["correct"] for e in ours["entries"]}
    for n, a in sorted(after.items()):
        if a["correct"] is None:
            continue
        mine = sum(1 for p in a["picks"] if p in our_nums)
        if mine != a["correct"]:
            problems += 1
            print(f"  total   #{n} {a['nick']}: his {a['correct']}, ours {mine}"
                  + (f" (was {our_correct[n]} on the card we parsed)" if our_correct.get(n) not in (None, mine) else ""))

    print("no differences" if not problems else f"{problems} difference(s)")


def _person(nick, name):
    """Row identity on Mike's suicide sheet: (nickname, real name), both stripped.

    Never the entry number. His graded sheet moves each week's eliminated
    entries to the top and renumbers them 1..N in column A, so those numbers
    collide with entries still alive (CLAUDE.md, Data Pipeline).
    """
    return (str(nick or "").strip(), str(name or "").strip())


def _is_red(cell):
    return bool(cell.fill.fill_type) and str(cell.fill.fgColor.rgb or "").upper().endswith("FF0000")


def _pair_unique(his_left, ours_left, his_val, our_val, pairs):
    """Pair leftover rows whose value is present exactly once on each side."""
    his_groups, our_groups = {}, {}
    for h in his_left:
        his_groups.setdefault(his_val(h), []).append(h)
    for o in ours_left:
        our_groups.setdefault(our_val(o), []).append(o)
    for v, hs in his_groups.items():
        os_ = our_groups.get(v, [])
        if v and len(hs) == 1 and len(os_) == 1:
            pairs.append((hs[0], os_[0]))
            his_left.remove(hs[0])
            ours_left.remove(os_[0])


def read_suicide_sheet(workbook):
    """Mike's graded suicide sheet -> one dict per row: key, sheet_no, picks, red.

    Week keys are strings, to match the JSON we compare and merge against.
    """
    import openpyxl
    sys.path.insert(0, str(ROOT / "scripts"))
    from parse_survivor import FIRST_WEEK_COL, MAX_WEEK_COL, normalise

    ws = openpyxl.load_workbook(workbook, data_only=True)["Sheet1"]
    his_rows = []
    for row in ws.iter_rows(min_row=3, max_row=ws.max_row):
        if row[0].value is None and not row[2].value:
            continue
        picks, red = {}, set()
        for col in range(FIRST_WEEK_COL, MAX_WEEK_COL + 1):
            wk = str(col - FIRST_WEEK_COL + 1)
            team, _ = normalise(row[col - 1].value)
            if team:
                picks[wk] = team
                if _is_red(row[col - 1]):
                    red.add(wk)
        his_rows.append({"key": _person(row[2].value, row[1].value), "sheet_no": row[0].value,
                         "picks": picks, "red": red})
    return his_rows


def pair_suicide_rows(his_rows, our_entries):
    """-> (pairs, his_left, ours_left), pairing on person, never on entry number.

    One pairing, used by both the check and the merge: two of them would be two
    chances to pair a row differently, and the merge would then write a pick
    onto an entry the check said was someone else's.
    """
    ours_by_key, his_by_key = {}, {}
    for e in our_entries:
        ours_by_key.setdefault(_person(e["nick"], e["name"]), []).append(e)
    for h in his_rows:
        his_by_key.setdefault(h["key"], []).append(h)

    # One person can hold several entries under the same nickname and name, so
    # each key is a multiset: identical pick histories pair first, then
    # whatever is left pairs in sheet order.
    pairs, his_left, ours_left = [], [], []
    for key in his_by_key.keys() | ours_by_key.keys():
        hs, os_ = list(his_by_key.get(key, [])), list(ours_by_key.get(key, []))
        for h in list(hs):
            same = next((o for o in os_ if o["picks"] == h["picks"]), None)
            if same:
                pairs.append((h, same))
                hs.remove(h)
                os_.remove(same)
        pairs.extend(zip(hs, os_))
        his_left.extend(hs[len(os_):])
        ours_left.extend(os_[len(hs):])

    # A nickname he retyped (AP8 -> AR8) leaves one unpaired row on each side.
    # Pair those on a real name, or failing that a pick history, that is
    # unique on both sides, and report a rename instead of a removal plus an add.
    _pair_unique(his_left, ours_left, lambda h: h["key"][1], lambda o: _person("", o["name"])[1], pairs)
    _pair_unique(his_left, ours_left, lambda h: json.dumps(h["picks"], sort_keys=True) if h["picks"] else "",
                 lambda o: json.dumps(o["picks"], sort_keys=True) if o["picks"] else "", pairs)
    return pairs, his_left, ours_left


def check_survivor(year, workbook, results=None):
    print("\n=== check: Mike's suicide sheet vs ours ===")
    his_rows = read_suicide_sheet(workbook)
    pairs, his_left, ours_left = pair_suicide_rows(
        his_rows, load(DATA / f"survivor-{year}.json")["entries"])

    problems = 0
    for h, o in sorted(pairs, key=lambda p: (p[1]["entry"] is None, p[1]["entry"] or 0)):
        label = f"#{o['entry']} {str(o['nick'] or '').strip()}"
        if _person(o["nick"], o["name"]) != h["key"]:
            problems += 1
            print(f"  renamed {label}: now {h['key'][0]!r}" + (f" / {h['key'][1]!r}" if h["key"][1] else ""))
        mine, his = o["picks"], h["picks"]
        for wk in sorted(mine.keys() & his.keys(), key=int):
            if mine[wk] != his[wk]:
                problems += 1
                print(f"  changed {label} week {wk}: {mine[wk]} -> {his[wk]}")
        for wk in sorted(mine.keys() - his.keys(), key=int):
            problems += 1
            print(f"  cleared {label} week {wk}: {mine[wk]} no longer on his sheet")
        for wk in sorted(his.keys() - mine.keys(), key=int):
            problems += 1
            print(f"  new     {label} week {wk}: {his[wk]} on his sheet only")
    for h in his_left:
        problems += 1
        print(f"  added   {h['key'][0]} (his row #{h['sheet_no']}): {h['picks']}")
    for o in ours_left:
        problems += 1
        print(f"  removed #{o['entry']} {o['nick']}")
    print("no pick differences" if not problems else f"{problems} pick difference(s)")

    # His red fills are his eliminations. Compare them with the losers we
    # graded, but only in weeks our grade has as final: his sheet can carry a
    # highlighted cell before the games are played, and that is not a result.
    for wk, w in sorted(((results or {}).get("weeks") or {}).items(), key=lambda x: int(x[0])):
        if w["pending"]:
            print(f"week {wk}: games still to play, red fills not compared")
            continue
        picked = [h for h in his_rows if wk in h["picks"]]
        red = [h for h in picked if wk in h["red"]]
        if not red:
            print(f"week {wk}: no red fills on his sheet, nothing to compare (we graded {w['lost']} lost)")
            continue
        his_out = Counter((h["key"], h["picks"][wk]) for h in red)
        our_out = Counter((_person(l["nick"], l["name"]), l["team"]) for l in w["losers"])
        if not his_out - our_out and his_out != our_out:
            # He marks early losers (a Thursday game) before mailing the sheet
            # that week, so a subset of ours is a sheet graded partway.
            print(f"week {wk}: his sheet is marked partway -- all {len(red)} of his red cells are "
                  f"among our {w['lost']} lost; the rest are not marked yet")
            continue
        verdict = "agree" if his_out == our_out else "DISAGREE"
        print(f"week {wk}: his {len(red)} out / {len(picked) - len(red)} alive; "
              f"ours {w['lost']} lost / {w['won']} won" + (f" / {w['tie']} tied" if w["tie"] else "")
              + f" -- {verdict}")
        for (key, team), n in sorted((his_out - our_out).items()):
            print(f"  out on his sheet only: {key[0]} ({team})" + (f" x{n}" if n > 1 else ""))
        for (key, team), n in sorted((our_out - his_out).items()):
            print(f"  lost in our grade only: {key[0]} ({team})" + (f" x{n}" if n > 1 else ""))


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[1])
    ap.add_argument("year", type=int)
    ap.add_argument("week", type=int, nargs="?", help="default: every week with parsed cards")
    ap.add_argument("--check", metavar="WEEKLY_PICKS_XLSX")
    ap.add_argument("--check-suicide", metavar="SUICIDE_XLSX")
    args = ap.parse_args()

    by_pair, by_team = schedule_index(args.year)
    weeks = [args.week] if args.week else sorted(
        int(p.stem.rsplit("-w", 1)[1]) for p in (DATA / "raw").glob(f"entries-{args.year}-w*.json"))

    graded = {w: grade_pickem(args.year, w, by_pair) for w in weeks}
    survivor = grade_survivor(args.year, by_team)

    if args.check:
        if len(weeks) != 1:
            sys.exit("--check needs a single week")
        check_pickem(args.year, weeks[0], graded[weeks[0]], args.check)
    if args.check_suicide:
        check_survivor(args.year, args.check_suicide, survivor)


if __name__ == "__main__":
    main()
